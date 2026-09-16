"""Word (.docx) extraction. Paragraphs are chunked; tables keep column labels like Excel."""

from pathlib import Path

import pandas as pd
from docx import Document

from app.services.ingestion.chunking import chunk_text
from app.services.ingestion.errors import IngestionError
from app.services.ingestion.tabular import records_from_frame
from app.services.ingestion.text_clean import clean_extracted_text


def extract_docx_chunks(path: Path, *, filename: str, category: str, region: str | None) -> tuple[int | None, list[dict]]:
    try:
        document = Document(path)
    except Exception as exc:
        raise IngestionError(f"Could not read Word document: {exc}") from exc

    records: list[dict] = []
    paragraphs = []
    for paragraph in document.paragraphs:
        cleaned = clean_extracted_text(paragraph.text or "")
        if cleaned:
            paragraphs.append(cleaned)

    body = clean_extracted_text("\n".join(paragraphs))
    for piece in chunk_text(body):
        records.append(
            {
                "content": piece,
                "category": category,
                "region": region,
                "filename": filename,
                "source_page": None,
                "source_row": None,
                "sheet_name": None,
                "section": "body",
                "chunk_index": len(records),
            }
        )

    for table_index, table in enumerate(document.tables, start=1):
        rows: list[list[str]] = []
        for row in table.rows:
            rows.append([(cell.text or "").strip() for cell in row.cells])
        if len(rows) < 2:
            continue
        header, *data_rows = rows
        if not any(header) or not any(any(cell for cell in row) for row in data_rows):
            continue
        frame = pd.DataFrame(data_rows, columns=header)
        records.extend(
            records_from_frame(
                frame,
                filename=filename,
                category=category,
                region=region,
                sheet_name=f"Table {table_index}",
                start_index=len(records),
            )
        )

    if not records:
        raise IngestionError("No extractable text was found in this Word document.")
    return None, records
