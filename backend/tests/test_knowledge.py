from io import BytesIO

import fitz
import pytest
from docx import Document
from openpyxl import Workbook

from app.auth.security import hash_password
from app.database.connection import document_chunks_collection, users_collection
from app.models.user import Role, UserStatus, new_user_document


def _pdf_bytes(text: str = "Sample leave policy: employees may request annual leave through HR.") -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), text)
    buffer = BytesIO()
    document.save(buffer)
    document.close()
    return buffer.getvalue()


def _xlsx_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Southern"
    sheet.append(["Office", "Hours"])
    sheet.append(["Galle", "08:00-16:00 weekday sample hours"])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("Sample leave policy: employees may request annual leave through HR.")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Office"
    table.rows[0].cells[1].text = "Hours"
    table.rows[1].cells[0].text = "Matara"
    table.rows[1].cells[1].text = "09:00-16:30"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


async def _insert_user(*, email: str, password: str, role: Role):
    document = new_user_document(
        name="Test",
        email=email,
        password_hash=hash_password(password),
        role=role,
        region="SOUTHERN" if role == Role.REGIONAL else None,
        status=UserStatus.ACTIVE,
    )
    result = await users_collection().insert_one(document)
    return str(result.inserted_id)


async def _login(client, email: str, password: str) -> str:
    response = await client.post("/auth/login", json={"email": email, "password": password})
    assert response.status_code == 200
    return response.json()["access_token"]


@pytest.mark.asyncio
async def test_normal_user_cannot_upload_knowledge(client):
    await _insert_user(email="normal@example.com", password="Password123", role=Role.NORMAL)
    token = await _login(client, "normal@example.com", "Password123")
    response = await client.post(
        "/admin/knowledge/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={"file": ("policy.pdf", _pdf_bytes(), "application/pdf")},
        data={"category": "GENERAL", "description": "test"},
    )
    assert response.status_code == 403


@pytest.mark.asyncio
async def test_admin_upload_pdf_stays_pending_until_approved(client):
    await _insert_user(email="admin1@example.com", password="AdminPass123", role=Role.ADMIN)
    token = await _login(client, "admin1@example.com", "AdminPass123")
    headers = {"Authorization": f"Bearer {token}"}

    upload = await client.post(
        "/admin/knowledge/upload",
        headers=headers,
        files={"file": ("hr_policy.pdf", _pdf_bytes(), "application/pdf")},
        data={"category": "GENERAL", "description": "Fake HR sample"},
    )
    assert upload.status_code == 201
    body = upload.json()
    assert body["status"] == "PENDING"
    assert body["chunk_count"] >= 1
    assert body["file_type"] == "pdf"
    document_id = body["id"]

    chunks = await client.get(f"/admin/knowledge/{document_id}/chunks", headers=headers)
    assert chunks.status_code == 200
    first = chunks.json()[0]
    assert first["source_page"] == 1
    assert first["status"] == "PENDING"
    assert "leave" in first["content"].lower()

    from app.auth.rbac import CurrentUser, chunk_is_visible_to

    super_user = CurrentUser(
        {
            "_id": "000000000000000000000099",
            "name": "Super",
            "email": "super@example.com",
            "role": Role.SUPER.value,
            "region": None,
            "status": UserStatus.ACTIVE.value,
        }
    )
    stored = await document_chunks_collection().find_one({"document_id": document_id})
    assert chunk_is_visible_to(super_user, stored) is False

    approve = await client.post(f"/admin/knowledge/{document_id}/approve", headers=headers)
    assert approve.status_code == 200
    assert approve.json()["status"] == "APPROVED"
    stored = await document_chunks_collection().find_one({"document_id": document_id})
    assert chunk_is_visible_to(super_user, stored) is True


@pytest.mark.asyncio
async def test_regional_upload_requires_region(client):
    await _insert_user(email="admin1@example.com", password="AdminPass123", role=Role.ADMIN)
    token = await _login(client, "admin1@example.com", "AdminPass123")
    response = await client.post(
        "/admin/knowledge/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={"file": ("south.pdf", _pdf_bytes(), "application/pdf")},
        data={"category": "REGIONAL", "description": "south"},
    )
    assert response.status_code == 400


@pytest.mark.asyncio
async def test_admin_upload_csv_keeps_column_labels(client):
    await _insert_user(email="admin1@example.com", password="AdminPass123", role=Role.ADMIN)
    token = await _login(client, "admin1@example.com", "AdminPass123")
    headers = {"Authorization": f"Bearer {token}"}
    csv_body = "Topic,Detail\nAnnual leave,Employees may request annual leave through HR\n"

    upload = await client.post(
        "/admin/knowledge/upload",
        headers=headers,
        files={"file": ("leave_faq.csv", csv_body.encode("utf-8"), "text/csv")},
        data={"category": "GENERAL", "description": "Fake FAQ table"},
    )
    assert upload.status_code == 201
    body = upload.json()
    assert body["status"] == "PENDING"
    assert body["file_type"] == "csv"
    assert body["chunk_count"] >= 1

    chunks = await client.get(f"/admin/knowledge/{body['id']}/chunks", headers=headers)
    content = chunks.json()[0]["content"].lower()
    assert "topic:" in content
    assert "annual leave" in content
    assert chunks.json()[0]["source_row"] == 2


@pytest.mark.asyncio
async def test_admin_upload_xlsx_preserves_sheet_name(client):
    await _insert_user(email="admin1@example.com", password="AdminPass123", role=Role.ADMIN)
    token = await _login(client, "admin1@example.com", "AdminPass123")
    headers = {"Authorization": f"Bearer {token}"}

    upload = await client.post(
        "/admin/knowledge/upload",
        headers=headers,
        files={
            "file": (
                "hours.xlsx",
                _xlsx_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={"category": "REGIONAL", "region": "SOUTHERN", "description": "Fake hours"},
    )
    assert upload.status_code == 201
    body = upload.json()
    assert body["file_type"] == "xlsx"
    assert body["region"] == "SOUTHERN"
    chunks = await client.get(f"/admin/knowledge/{body['id']}/chunks", headers=headers)
    first = chunks.json()[0]
    assert first["sheet_name"] == "Southern"
    assert "galle" in first["content"].lower()


@pytest.mark.asyncio
async def test_super_user_can_upload_and_approve_but_cannot_delete(client):
    await _insert_user(email="super@example.com", password="Password123", role=Role.SUPER)
    token = await _login(client, "super@example.com", "Password123")
    headers = {"Authorization": f"Bearer {token}"}

    upload = await client.post(
        "/admin/knowledge/upload",
        headers=headers,
        files={"file": ("hr_policy.pdf", _pdf_bytes(), "application/pdf")},
        data={"category": "GENERAL", "description": "Super User validation sample"},
    )
    assert upload.status_code == 201
    body = upload.json()
    assert body["status"] == "PENDING"
    document_id = body["id"]

    approve = await client.post(f"/admin/knowledge/{document_id}/approve", headers=headers)
    assert approve.status_code == 200
    assert approve.json()["status"] == "APPROVED"

    listed = await client.get("/admin/knowledge", headers=headers)
    assert listed.status_code == 200
    assert any(row["id"] == document_id for row in listed.json())

    deleted = await client.delete(f"/admin/knowledge/{document_id}", headers=headers)
    assert deleted.status_code == 403


@pytest.mark.asyncio
async def test_super_user_upload_docx_stays_pending_until_approved(client):
    await _insert_user(email="super@example.com", password="Password123", role=Role.SUPER)
    token = await _login(client, "super@example.com", "Password123")
    headers = {"Authorization": f"Bearer {token}"}

    upload = await client.post(
        "/admin/knowledge/upload",
        headers=headers,
        files={
            "file": (
                "leave_policy.docx",
                _docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        data={"category": "GENERAL", "description": "Word sample"},
    )
    assert upload.status_code == 201
    body = upload.json()
    assert body["status"] == "PENDING"
    assert body["file_type"] == "docx"
    assert body["chunk_count"] >= 2

    chunks = await client.get(f"/admin/knowledge/{body['id']}/chunks", headers=headers)
    assert chunks.status_code == 200
    contents = " ".join(item["content"] for item in chunks.json()).lower()
    assert "annual leave" in contents
    assert "matara" in contents
    assert any(item.get("sheet_name") == "Table 1" for item in chunks.json())

    approve = await client.post(f"/admin/knowledge/{body['id']}/approve", headers=headers)
    assert approve.status_code == 200
    assert approve.json()["status"] == "APPROVED"
